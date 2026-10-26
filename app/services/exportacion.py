"""Exportación del recorte visible a Excel y a PDF.

Principio que gobierna los dos formatos: **un archivo exportado tiene que poder
explicarse solo**. Seis meses después nadie recuerda con qué filtros se generó una
planilla, y una cifra sin sus criterios no se puede reproducir ni discutir. Por eso
los dos formatos incluyen, de manera destacada, qué filtros se aplicaron, cuántos
ensayos abarcan, quién exportó y cuándo.

Reparto de responsabilidades:

* Excel es para **seguir trabajando**: se entrega en dos hojas de datos, una ancha
  (un ensayo por fila, para leer) y una larga (un producto por fila, para tablas
  dinámicas). Convertir de larga a ancha en Excel es tedioso; al revés, imposible.
* PDF es para **circular**: resumen, gráficos y las primeras filas. No pretende
  llevar el dato completo; para eso está el Excel.

Los dos consumen exactamente las mismas funciones de `analitica`, así que un número
del informe no puede diferir del que muestra la pantalla.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.orm import joinedload, selectinload

from ..extensions import db
from ..models import DetalleMezcla, Mezcla, Tratamiento
from ..security import registrar
from . import analitica, consultas

# Tope de filas. Un histórico de decenas de miles de ensayos exportado entero
# consume memoria del servidor y produce un archivo que nadie abre. Se corta, y el
# archivo lo dice en la hoja de criterios en lugar de mentir por omisión.
MAXIMO_FILAS = 20_000
FILAS_EN_PDF = 40

COLUMNAS_ANCHA = [
    ("N° tratamiento", 16),
    ("Fecha", 12),
    ("Cliente", 30),
    ("Cultivo", 14),
    ("Tipo", 18),
    ("Target", 20),
    ("Productos de la mezcla", 52),
    ("Sin mapear", 12),
    ("pH agua", 10),
    ("Dureza agua", 13),
    ("pH mezcla", 11),
    ("Obs. 5 min", 16),
    ("Obs. 10 min", 16),
    ("Observación final", 22),
    ("Origen", 13),
    ("Estado", 16),
]

COLUMNAS_LARGA = [
    ("N° tratamiento", 16),
    ("Fecha", 12),
    ("Cliente", 30),
    ("Cultivo", 14),
    ("Orden", 8),
    ("Nombre en la planilla", 34),
    ("Producto del maestro", 30),
    ("Principio activo", 30),
    ("Categoría", 16),
    ("Estado de normalización", 22),
    ("Dosis", 10),
    ("Unidad", 10),
]


@dataclass
class Contexto:
    """Quién exportó, cuándo y con qué criterios. Es la portada de los dos formatos."""

    usuario: str
    momento: datetime
    criterios: list[str]
    total: int
    exportadas: int

    @property
    def recortado(self) -> bool:
        return self.exportadas < self.total


def _contexto(filtros: consultas.Filtros, usuario, total: int, exportadas: int) -> Contexto:
    criterios = filtros.descripcion() or ["sin filtros: todo el repositorio"]
    return Contexto(
        usuario=getattr(usuario, "nombre", "—"),
        momento=datetime.now(),
        criterios=criterios,
        total=total,
        exportadas=exportadas,
    )


def _tratamientos(filtros: consultas.Filtros, limite: int = MAXIMO_FILAS) -> list[Tratamiento]:
    """El recorte, con mezclas y resultado ya cargados.

    Se cargan las relaciones por adelantado porque el bucle de escritura recorre
    todos los productos de cada ensayo: sin esto, exportar mil ensayos serían miles
    de consultas sueltas.
    """
    consulta = (
        select(Tratamiento)
        .where(Tratamiento.id.in_(consultas.subconsulta_ids(filtros)))
        .options(
            joinedload(Tratamiento.cliente),
            joinedload(Tratamiento.resultado),
            selectinload(Tratamiento.mezclas)
            .selectinload(Mezcla.detalles)
            .joinedload(DetalleMezcla.producto),
        )
        .order_by(Tratamiento.fecha.desc().nullslast(), Tratamiento.id.desc())
        .limit(limite)
    )
    return list(db.session.scalars(consulta).unique())


def _detalles(tratamiento: Tratamiento) -> list[DetalleMezcla]:
    return [d for m in tratamiento.mezclas for d in m.detalles]


def _texto_mezcla(tratamiento: Tratamiento) -> str:
    """La mezcla como la leería una persona: el nombre resuelto, o el crudo si no lo hay."""
    nombres = [
        (d.producto.nombre_comercial if d.producto else d.nombre_original)
        for d in _detalles(tratamiento)
    ]
    return " + ".join(nombres)


def _sin_mapear(tratamiento: Tratamiento) -> int:
    from ..models import EstadoNormalizacion

    return sum(
        1
        for d in _detalles(tratamiento)
        if d.estado_normalizacion is EstadoNormalizacion.PENDIENTE
    )


def _numero(valor) -> float | None:
    return float(valor) if valor is not None else None


# -------------------------------------------------------------------- Excel


def libro_excel(filtros: consultas.Filtros, usuario) -> tuple[BytesIO, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    total = analitica.resumen(filtros).tratamientos
    tratamientos = _tratamientos(filtros)
    contexto = _contexto(filtros, usuario, total, len(tratamientos))

    libro = Workbook()
    encabezado = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="0D5C4A")

    def escribir_cabecera(hoja, columnas):
        for indice, (titulo, ancho) in enumerate(columnas, start=1):
            celda = hoja.cell(row=1, column=indice, value=titulo)
            celda.font = encabezado
            celda.fill = fondo
            celda.alignment = Alignment(vertical="center")
            hoja.column_dimensions[get_column_letter(indice)].width = ancho
        # Fila fija y autofiltro: el que recibe el archivo espera poder filtrar.
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = (
            f"A1:{get_column_letter(len(columnas))}{max(hoja.max_row, 1)}"
        )

    # --- Hoja 1: criterios. Va primera a propósito: es lo que hace reproducible al resto.
    criterios = libro.active
    criterios.title = "Criterios"
    criterios["A1"] = "FormData — exportación de tratamientos"
    criterios["A1"].font = Font(bold=True, size=14)
    filas_criterios = [
        ("Generado por", contexto.usuario),
        ("Fecha y hora", contexto.momento.strftime("%d/%m/%Y %H:%M")),
        ("Ensayos que cumplen los criterios", contexto.total),
        ("Ensayos incluidos en este archivo", contexto.exportadas),
    ]
    if contexto.recortado:
        filas_criterios.append(
            ("Atención", f"El archivo se recortó al máximo de {MAXIMO_FILAS} filas. Afiná los filtros.")
        )
    for indice, (etiqueta, valor) in enumerate(filas_criterios, start=3):
        criterios.cell(row=indice, column=1, value=etiqueta).font = Font(bold=True)
        criterios.cell(row=indice, column=2, value=valor)

    fila = len(filas_criterios) + 4
    criterios.cell(row=fila, column=1, value="Filtros aplicados").font = Font(bold=True)
    for indice, criterio in enumerate(contexto.criterios, start=fila + 1):
        criterios.cell(row=indice, column=2, value=criterio)
    criterios.column_dimensions["A"].width = 34
    criterios.column_dimensions["B"].width = 62

    # --- Hoja 2: resumen e indicadores
    resumen = libro.create_sheet("Resumen")
    datos = analitica.resumen(filtros)
    promedios = analitica.promedios(filtros)
    resumen["A1"] = "Cifras del recorte"
    resumen["A1"].font = Font(bold=True, size=12)
    for indice, (etiqueta, valor) in enumerate(
        [
            ("Ensayos", datos.tratamientos),
            ("Clientes distintos", datos.clientes),
            ("Cultivos distintos", datos.cultivos),
            ("Productos distintos usados", datos.productos),
            ("Filas de producto", datos.filas_producto),
            ("Filas sin mapear", datos.filas_sin_mapear),
            ("Normalización del recorte (%)", datos.porcentaje_normalizado),
            (f"pH de agua promedio (n={promedios.medidos_ph_agua})", promedios.ph_agua),
            (f"pH de mezcla promedio (n={promedios.medidos_ph_mezcla})", promedios.ph_mezcla),
            (f"Dureza de agua promedio (n={promedios.medidos_dureza})", promedios.dureza_agua),
        ],
        start=3,
    ):
        resumen.cell(row=indice, column=1, value=etiqueta).font = Font(bold=True)
        resumen.cell(row=indice, column=2, value=valor)

    fila = 15
    for serie in (
        analitica.por_cultivo(filtros),
        analitica.productos_mas_usados(filtros),
        analitica.observaciones_de_ensayo(filtros),
    ):
        resumen.cell(row=fila, column=1, value=serie.titulo).font = Font(bold=True, size=12)
        fila += 1
        for punto in serie.puntos:
            resumen.cell(row=fila, column=1, value=punto.etiqueta)
            resumen.cell(row=fila, column=2, value=punto.valor)
            fila += 1
        fila += 1
    resumen.column_dimensions["A"].width = 40
    resumen.column_dimensions["B"].width = 16

    # --- Hoja 3: una fila por ensayo (formato ancho, para leer)
    hoja = libro.create_sheet("Tratamientos")
    for tratamiento in tratamientos:
        resultado = tratamiento.resultado
        hoja.append(
            [
                tratamiento.numero_tratamiento,
                tratamiento.fecha,
                tratamiento.cliente.nombre if tratamiento.cliente else None,
                tratamiento.cultivo,
                tratamiento.tipo_tratamiento,
                tratamiento.target,
                _texto_mezcla(tratamiento),
                _sin_mapear(tratamiento),
                _numero(resultado.ph_agua) if resultado else None,
                _numero(resultado.dureza_agua) if resultado else None,
                _numero(resultado.ph_mezcla) if resultado else None,
                resultado.observacion_cinco_min if resultado else None,
                resultado.observacion_diez_min if resultado else None,
                resultado.estado_resultado if resultado else None,
                tratamiento.origen.value,
                tratamiento.estado_flujo.etiqueta,
            ]
        )
    hoja.insert_rows(1)
    escribir_cabecera(hoja, COLUMNAS_ANCHA)
    for celda in hoja["B"][1:]:
        celda.number_format = "DD/MM/YYYY"

    # --- Hoja 4: una fila por producto (formato largo, para tablas dinámicas)
    larga = libro.create_sheet("Productos por ensayo")
    for tratamiento in tratamientos:
        for detalle in _detalles(tratamiento):
            producto = detalle.producto
            larga.append(
                [
                    tratamiento.numero_tratamiento,
                    tratamiento.fecha,
                    tratamiento.cliente.nombre if tratamiento.cliente else None,
                    tratamiento.cultivo,
                    detalle.orden_producto,
                    detalle.nombre_original,
                    producto.nombre_comercial if producto else None,
                    producto.activo if producto else None,
                    producto.categoria if producto else None,
                    detalle.estado_normalizacion.value,
                    _numero(detalle.dosis),
                    detalle.unidad,
                ]
            )
    larga.insert_rows(1)
    escribir_cabecera(larga, COLUMNAS_LARGA)
    for celda in larga["B"][1:]:
        celda.number_format = "DD/MM/YYYY"

    flujo = BytesIO()
    libro.save(flujo)
    flujo.seek(0)

    registrar(
        "exportacion_excel",
        entidad="tratamiento",
        usuario_id=getattr(usuario, "id", None),
        datos={"filas": len(tratamientos), "criterios": contexto.criterios},
    )
    return flujo, _nombre_archivo(contexto, "xlsx")


# ---------------------------------------------------------------------- PDF


def informe_pdf(filtros: consultas.Filtros, usuario) -> tuple[BytesIO, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    verde = colors.HexColor("#0d5c4a")
    tinta = colors.HexColor("#12211d")
    linea = colors.HexColor("#d7ded8")
    papel = colors.HexColor("#eef1ec")

    total = analitica.resumen(filtros).tratamientos
    tratamientos = _tratamientos(filtros, limite=FILAS_EN_PDF)
    contexto = _contexto(filtros, usuario, total, len(tratamientos))

    flujo = BytesIO()
    documento = SimpleDocTemplate(
        flujo,
        pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=16 * mm,
        title="FormData — informe de tratamientos",
        author="FormData",
    )

    base = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle("titulo", parent=base["Title"], textColor=verde, fontSize=19,
                                   spaceAfter=2, alignment=0)
    estilo_bajada = ParagraphStyle("bajada", parent=base["Normal"], textColor=tinta, fontSize=9.5,
                                   spaceAfter=10)
    estilo_h2 = ParagraphStyle("h2", parent=base["Heading2"], textColor=verde, fontSize=12,
                               spaceBefore=10, spaceAfter=5)
    estilo_nota = ParagraphStyle("nota", parent=base["Normal"], fontSize=8,
                                 textColor=colors.HexColor("#7e8d87"))
    estilo_celda = ParagraphStyle("celda", parent=base["Normal"], fontSize=7.4, leading=9)

    elementos = [
        Paragraph("Informe de tratamientos", estilo_titulo),
        Paragraph(
            f"Generado por {contexto.usuario} el "
            f"{contexto.momento.strftime('%d/%m/%Y a las %H:%M')}",
            estilo_bajada,
        ),
    ]

    # Criterios. Van arriba de todo: sin ellos las cifras de abajo no se pueden repetir.
    elementos.append(Paragraph("Criterios aplicados", estilo_h2))
    tabla_criterios = Table(
        [[Paragraph("· " + c, estilo_celda)] for c in contexto.criterios],
        colWidths=[documento.width],
    )
    tabla_criterios.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), papel),
                ("BOX", (0, 0), (-1, -1), 0.5, linea),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elementos.append(tabla_criterios)

    # Cifras
    datos = analitica.resumen(filtros)
    promedios = analitica.promedios(filtros)
    elementos.append(Paragraph("Cifras del recorte", estilo_h2))
    cifras = [
        ["Ensayos", "Clientes", "Cultivos", "Productos", "Normalización", "pH mezcla prom."],
        [
            f"{datos.tratamientos}",
            f"{datos.clientes}",
            f"{datos.cultivos}",
            f"{datos.productos}",
            f"{datos.porcentaje_normalizado} %",
            f"{promedios.ph_mezcla if promedios.ph_mezcla is not None else '—'} "
            f"(n={promedios.medidos_ph_mezcla})",
        ],
    ]
    tabla_cifras = Table(cifras, colWidths=[documento.width / 6] * 6)
    tabla_cifras.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, 0), 7.5),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#7e8d87")),
                ("FONTSIZE", (0, 1), (-1, 1), 14),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, 1), (-1, 1), verde),
                ("BOX", (0, 0), (-1, -1), 0.5, linea),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, linea),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elementos.append(tabla_cifras)

    if not datos.confiable:
        elementos.append(Spacer(1, 5))
        elementos.append(
            Paragraph(
                f"Atención: {datos.filas_sin_mapear} fila(s) de producto de este recorte "
                "todavía no resolvieron contra el maestro. Los rankings de producto de este "
                "informe subestiman esos casos.",
                estilo_nota,
            )
        )

    # Gráficos: las mismas series de la pantalla, dibujadas con barras de tabla.
    for serie in (
        analitica.por_cultivo(filtros, limite=6),
        analitica.productos_mas_usados(filtros, limite=6),
        analitica.observaciones_de_ensayo(filtros, limite=6),
    ):
        if not serie.hay_datos:
            continue
        elementos.append(
            KeepTogether([Paragraph(serie.titulo, estilo_h2), _barras_pdf(serie, documento.width, verde, papel, linea)])
        )

    # Detalle
    elementos.append(PageBreak())
    elementos.append(Paragraph("Detalle de ensayos", estilo_h2))
    if contexto.recortado:
        elementos.append(
            Paragraph(
                f"Se listan los {contexto.exportadas} ensayos más recientes de "
                f"{contexto.total} que cumplen los criterios. El listado completo está en "
                "la exportación a Excel.",
                estilo_nota,
            )
        )
        elementos.append(Spacer(1, 5))

    cabeceras = ["N°", "Fecha", "Cliente", "Cultivo", "Mezcla", "pH mezcla", "Observación"]
    filas = [[Paragraph(f"<b>{c}</b>", estilo_celda) for c in cabeceras]]
    for tratamiento in tratamientos:
        resultado = tratamiento.resultado
        filas.append(
            [
                Paragraph(tratamiento.numero_tratamiento, estilo_celda),
                Paragraph(
                    tratamiento.fecha.strftime("%d/%m/%Y") if tratamiento.fecha else "—",
                    estilo_celda,
                ),
                Paragraph(tratamiento.cliente.nombre if tratamiento.cliente else "—", estilo_celda),
                Paragraph(tratamiento.cultivo or "—", estilo_celda),
                Paragraph(_texto_mezcla(tratamiento) or "—", estilo_celda),
                Paragraph(
                    str(_numero(resultado.ph_mezcla)) if resultado and resultado.ph_mezcla is not None else "—",
                    estilo_celda,
                ),
                Paragraph(
                    (resultado.estado_resultado if resultado else None) or "—", estilo_celda
                ),
            ]
        )

    # La primera columna tiene que entrar «SOL-2026-0005» en un renglón: partir un
    # identificador en dos líneas lo vuelve imposible de buscar de un vistazo.
    anchos = [0.11, 0.08, 0.16, 0.09, 0.33, 0.08, 0.15]
    tabla = Table(filas, colWidths=[documento.width * a for a in anchos], repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), verde),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, linea),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                # Filas alternadas: con siete columnas angostas, seguir un renglón
                # de punta a punta sin guía es donde se cometen los errores de lectura.
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, papel]),
            ]
        )
    )
    elementos.append(tabla)

    documento.build(elementos, onLaterPages=_pie, onFirstPage=_pie)
    flujo.seek(0)

    registrar(
        "exportacion_pdf",
        entidad="tratamiento",
        usuario_id=getattr(usuario, "id", None),
        datos={"filas": len(tratamientos), "criterios": contexto.criterios},
    )
    return flujo, _nombre_archivo(contexto, "pdf")


def _barras_pdf(serie, ancho_total, verde, papel, linea):
    """Las barras del PDF, hechas con celdas de tabla coloreadas.

    Reutilizar el SVG habría exigido `svglib` solo para esto. Una tabla de dos
    columnas con la celda de la barra pintada al ancho proporcional da el mismo
    resultado con lo que reportlab ya trae.
    """
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, Table, TableStyle

    base = getSampleStyleSheet()
    estilo = ParagraphStyle("barra", parent=base["Normal"], fontSize=8, leading=10)

    ancho_etiqueta = ancho_total * 0.28
    ancho_pista = ancho_total * 0.62
    ancho_valor = ancho_total * 0.10

    filas, estilos = [], [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (1, 0), (1, -1), 0),
        ("RIGHTPADDING", (1, 0), (1, -1), 0),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"),
    ]

    for indice, punto in enumerate(serie.puntos):
        proporcion = serie.porcentaje(punto) / 100
        pintado = max(ancho_pista * proporcion, 0.5)
        # La barra es una tabla anidada de una celda con el ancho ya calculado.
        barra = Table([[""]], colWidths=[pintado], rowHeights=[7])
        barra.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), verde)]))
        interior = Table([[barra, ""]], colWidths=[pintado, max(ancho_pista - pintado, 0.1)])
        interior.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (1, 0), (1, 0), papel),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        filas.append([Paragraph(punto.etiqueta, estilo), interior, Paragraph(str(int(punto.valor)), estilo)])

    tabla = Table(filas, colWidths=[ancho_etiqueta, ancho_pista, ancho_valor])
    tabla.setStyle(TableStyle(estilos))
    return tabla


def _pie(canvas, documento):
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#7e8d87"))
    canvas.drawString(14 * mm, 9 * mm, "FormData · Laboratorios Bahnsa")
    canvas.drawRightString(
        documento.pagesize[0] - 14 * mm, 9 * mm, f"Página {canvas.getPageNumber()}"
    )
    canvas.restoreState()


def _nombre_archivo(contexto: Contexto, extension: str) -> str:
    return f"formdata-tratamientos-{contexto.momento:%Y%m%d-%H%M}.{extension}"
