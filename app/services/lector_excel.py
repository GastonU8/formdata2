"""Lectura del Excel de tratamientos.

Este módulo solo lee y valida: no toca la base. Devuelve filas limpias o errores
con número de fila y columna, para que la pantalla de importación pueda mostrar
exactamente qué corregir.

La tolerancia de los encabezados es deliberada. El archivo real trae cosas como
`Granulometría (>180 µm)22` — basura pegada de una edición vieja. Un validador
estricto rechazaría el archivo entero por eso; uno que normaliza el encabezado
lo lee igual y sigue.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Iterator

from openpyxl import load_workbook

HOJA_TRATAMIENTOS = "Tratamientos"
HOJA_PRODUCTOS = "Productos Normalizados"
HOJA_MAPEO = "Mapeo_Productos"
HOJA_MAESTRO = "Maestro_Productos"

SEPARADOR_MEZCLA = "+"


class ArchivoInvalido(Exception):
    """El archivo no tiene la forma mínima para poder procesarse."""

    def __init__(self, errores: list[str]):
        self.errores = errores
        super().__init__("; ".join(errores))


def _clave_encabezado(texto: Any) -> str:
    """Normaliza un encabezado para poder compararlo sin sorpresas."""
    if texto is None:
        return ""
    plano = unicodedata.normalize("NFKD", str(texto))
    plano = "".join(c for c in plano if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", plano).split())


# Encabezado normalizado -> nombre interno del campo.
COLUMNAS_TRATAMIENTO = {
    "fecha": "fecha",
    "n de tratamiento": "numero",
    "tratamiento": "mezcla_texto",
    "link": "enlace",
    "pais": "pais",
    "lab": "laboratorio",
    "cliente": "cliente",
    "nro cuenta": "numero_cuenta",
    "cliente softland": "cliente_softland",
    "ubicacion": "ubicacion",
    "clasificacion": "clasificacion",
    "cantidad ha": "cantidad_ha",
    "cultivo": "cultivo",
    "momento aplicacion": "momento_aplicacion",
    "target": "target",
    "tipo tratamiento": "tipo_tratamiento",
    "tasaaplicacion": "tasa_aplicacion",
    "ph agua": "ph_agua",
    "dureza agua": "dureza_agua",
    "aguacorregida": "agua_corregida",
    "ph aguacorregida": "ph_agua_corregida",
    "dureza aguacorregida": "dureza_agua_corregida",
    "ph mezcla": "ph_mezcla",
    "observacion 5 min": "observacion_cinco_min",
    "volum ml 5 min": "volumen_cinco_min",
    "observacion 10 min": "observacion_diez_min",
    "volum ml 10 min": "volumen_diez_min",
    "suspensibilidad n inversiones": "suspensibilidad",
    "dosis de compatibilizador": "dosis_compatibilizador",
    "observaciones": "observaciones",
}

# `Año` y `Mes` se ignoran: se derivan de `Fecha`. En el archivo real las filas
# vacías traen Año=1900 y Mes=1 por fórmulas arrastradas, así que además mienten.
COLUMNAS_IGNORADAS = {"ano", "anio", "mes"}

COLUMNAS_PRODUCTO = {
    "n de tratamiento": "numero",
    "producto": "nombre",
    "orden": "orden",
    # `Producto_Final` e `Id_Producto` se leen pero NO se usan para resolver:
    # en el archivo real la fórmula devuelve productos equivocados en lugar de
    # "SIN MAPEAR". El sistema vuelve a resolver desde el nombre crudo.
    "producto final": "producto_final_excel",
    "id producto": "id_producto_excel",
}

RE_MICRAS = re.compile(r">\s*(\d+)")
VERDADEROS = {"si", "sí", "s", "true", "verdadero", "1", "x"}
FALSOS = {"no", "n", "false", "falso", "0"}


@dataclass
class ErrorFila:
    fila: int
    columna: str
    mensaje: str

    def como_dict(self) -> dict:
        return {"fila": self.fila, "columna": self.columna, "mensaje": self.mensaje}


@dataclass
class Lectura:
    """Resultado de leer el archivo: qué se pudo leer y qué no."""

    tratamientos: list[dict] = field(default_factory=list)
    productos_por_tratamiento: dict[str, list[dict]] = field(default_factory=dict)
    mapeos: list[tuple[str, str]] = field(default_factory=list)
    maestro: list[dict] = field(default_factory=list)
    errores: list[ErrorFila] = field(default_factory=list)
    filas_leidas: int = 0
    filas_omitidas: int = 0


# --------------------------------------------------------------- conversores


def _texto(valor: Any, largo: int = 255) -> str | None:
    if valor is None:
        return None
    limpio = str(valor).strip()
    return limpio[:largo] or None


def _numero(valor: Any) -> float | None:
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("%", "")
    # El archivo trae cosas como "< 0.50": no es un número, es una observación.
    if not texto or not re.fullmatch(r"-?[\d.,]+", texto):
        return None
    texto = texto.replace(".", "").replace(",", ".") if texto.count(",") == 1 and texto.count(".") > 1 else texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _booleano(valor: Any) -> bool | None:
    if valor is None:
        return None
    if isinstance(valor, bool):
        return valor
    texto = str(valor).strip().lower()
    if texto in VERDADEROS:
        return True
    if texto in FALSOS:
        return False
    return None


def _fecha(valor: Any) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if valor is None:
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(valor).strip(), formato).date()
        except ValueError:
            continue
    return None


def _numero_tratamiento(valor: Any) -> str | None:
    """El número puede venir como 1990 o 1990.0 según cómo lo guardó Excel."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


def partir_mezcla(texto: str | None) -> list[str]:
    """Desglosa la columna `Tratamiento` usando `+` como separador.

    Es la misma regla que aplica la macro `Normalizar_Incremental` del archivo.
    Se usa solo como respaldo: si la hoja `Productos Normalizados` ya tiene el
    desglose, manda esa.
    """
    if not texto:
        return []
    return [parte.strip() for parte in str(texto).split(SEPARADOR_MEZCLA) if parte.strip()]


# ------------------------------------------------------------------- lectura


def _mapa_columnas(encabezados: tuple, alias: dict[str, str]) -> dict[str, int]:
    """Ubica cada campo conocido en su índice de columna."""
    mapa: dict[str, int] = {}
    for indice, bruto in enumerate(encabezados):
        clave = _clave_encabezado(bruto)
        if not clave or clave in COLUMNAS_IGNORADAS:
            continue
        if clave in alias:
            mapa.setdefault(alias[clave], indice)
            continue
        # Granulometría por rango de micras: el encabezado real puede traer
        # basura pegada, así que se identifica por el número dentro del ">".
        if clave.startswith("granulometria"):
            micras = RE_MICRAS.search(str(bruto))
            if micras and micras.group(1) in {"300", "180", "150", "106", "53"}:
                mapa.setdefault(f"granulometria_{micras.group(1)}", indice)
    return mapa


def _valor(fila: tuple, mapa: dict[str, int], campo: str) -> Any:
    indice = mapa.get(campo)
    if indice is None or indice >= len(fila):
        return None
    return fila[indice]


def leer(ruta: str) -> Lectura:
    """Lee el archivo completo. Levanta `ArchivoInvalido` si no sirve."""
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:  # openpyxl levanta varios tipos distintos
        raise ArchivoInvalido([f"No se pudo abrir el archivo: {error}"]) from error

    faltantes = [h for h in (HOJA_TRATAMIENTOS,) if h not in libro.sheetnames]
    if faltantes:
        raise ArchivoInvalido(
            [f"Falta la hoja «{h}». El archivo no parece el de tratamientos." for h in faltantes]
        )

    lectura = Lectura()
    _leer_mapeo(libro, lectura)
    _leer_maestro(libro, lectura)
    _leer_productos(libro, lectura)
    _leer_tratamientos(libro, lectura)
    libro.close()

    if not lectura.tratamientos:
        raise ArchivoInvalido(
            ["La hoja «Tratamientos» no tiene ninguna fila con N° de tratamiento."]
        )
    return lectura


def _filas(libro, hoja: str) -> Iterator[tuple[int, tuple]]:
    ws = libro[hoja]
    for numero, fila in enumerate(ws.iter_rows(values_only=True), start=1):
        yield numero, fila


def _leer_mapeo(libro, lectura: Lectura) -> None:
    if HOJA_MAPEO not in libro.sheetnames:
        return
    filas = _filas(libro, HOJA_MAPEO)
    next(filas, None)  # encabezado
    for _, fila in filas:
        original, final = (_texto(fila[0], 220) if fila else None), (
            _texto(fila[1], 220) if len(fila) > 1 else None
        )
        if original and final:
            lectura.mapeos.append((original, final))


def _leer_maestro(libro, lectura: Lectura) -> None:
    if HOJA_MAESTRO not in libro.sheetnames:
        return
    filas = _filas(libro, HOJA_MAESTRO)
    _, encabezados = next(filas, (0, ()))
    alias = {
        "id producto": "id_externo",
        "producto final": "nombre",
        "n registro": "numero_registro",
        "empresa": "empresa",
        "activos": "activos",
        "categoria": "categoria",
        "tipo de formulacion": "formulacion",
    }
    mapa = _mapa_columnas(encabezados, alias)
    for _, fila in filas:
        nombre = _texto(_valor(fila, mapa, "nombre"), 220)
        if not nombre:
            continue
        lectura.maestro.append(
            {
                "id_externo": _texto(_valor(fila, mapa, "id_externo"), 30),
                "nombre": nombre,
                "numero_registro": _texto(_valor(fila, mapa, "numero_registro"), 60),
                "empresa": _texto(_valor(fila, mapa, "empresa"), 120),
                "activos": _texto(_valor(fila, mapa, "activos"), 180),
                "categoria": _texto(_valor(fila, mapa, "categoria"), 120),
                "formulacion": _texto(_valor(fila, mapa, "formulacion"), 60),
            }
        )


def _leer_productos(libro, lectura: Lectura) -> None:
    if HOJA_PRODUCTOS not in libro.sheetnames:
        return
    filas = _filas(libro, HOJA_PRODUCTOS)
    _, encabezados = next(filas, (0, ()))
    mapa = _mapa_columnas(encabezados, COLUMNAS_PRODUCTO)

    for numero_fila, fila in filas:
        numero = _numero_tratamiento(_valor(fila, mapa, "numero"))
        nombre = _texto(_valor(fila, mapa, "nombre"), 220)
        if not numero or not nombre:
            continue
        orden = _numero(_valor(fila, mapa, "orden"))
        lectura.productos_por_tratamiento.setdefault(numero, []).append(
            {
                "nombre": nombre,
                "orden": int(orden) if orden else len(lectura.productos_por_tratamiento[numero]) + 1,
                "fila": numero_fila,
            }
        )


def _leer_tratamientos(libro, lectura: Lectura) -> None:
    filas = _filas(libro, HOJA_TRATAMIENTOS)
    _, encabezados = next(filas, (0, ()))
    mapa = _mapa_columnas(encabezados, COLUMNAS_TRATAMIENTO)

    esenciales = {"numero", "mezcla_texto"}
    ausentes = esenciales - set(mapa)
    if ausentes:
        raise ArchivoInvalido(
            [
                "La hoja «Tratamientos» no tiene la columna "
                + {"numero": "«N° de tratamiento»", "mezcla_texto": "«Tratamiento»"}[campo]
                for campo in sorted(ausentes)
            ]
        )

    vistos: dict[str, int] = {}

    for numero_fila, fila in filas:
        numero = _numero_tratamiento(_valor(fila, mapa, "numero"))

        # Sin número no hay tratamiento. En el archivo real hay miles de filas
        # así, con Año=1900 por fórmulas arrastradas hasta el fondo de la hoja.
        if not numero:
            if any(v is not None and str(v).strip() for v in fila):
                lectura.filas_omitidas += 1
            continue

        lectura.filas_leidas += 1

        if numero in vistos:
            lectura.errores.append(
                ErrorFila(
                    numero_fila, "N° de tratamiento",
                    f"El número {numero} ya aparece en la fila {vistos[numero]} del mismo archivo.",
                )
            )
            continue
        vistos[numero] = numero_fila

        fecha = _fecha(_valor(fila, mapa, "fecha"))
        if fecha is None and _valor(fila, mapa, "fecha") is not None:
            lectura.errores.append(
                ErrorFila(numero_fila, "Fecha", "No se pudo interpretar la fecha.")
            )

        productos = lectura.productos_por_tratamiento.get(numero, [])
        mezcla_texto = _texto(_valor(fila, mapa, "mezcla_texto"), 2000)
        if not productos:
            # Respaldo: desglosar la columna `Tratamiento` con el mismo criterio
            # que la macro del archivo.
            productos = [
                {"nombre": nombre, "orden": orden, "fila": numero_fila}
                for orden, nombre in enumerate(partir_mezcla(mezcla_texto), start=1)
            ]
            if not productos:
                lectura.errores.append(
                    ErrorFila(numero_fila, "Tratamiento", "El tratamiento no tiene productos.")
                )

        datos = {"numero": numero, "fila": numero_fila, "fecha": fecha, "productos": productos}
        for campo in (
            "mezcla_texto", "enlace", "pais", "laboratorio", "cliente", "numero_cuenta",
            "cliente_softland", "ubicacion", "clasificacion", "cultivo",
            "momento_aplicacion", "target", "tipo_tratamiento", "observaciones",
            "observacion_cinco_min", "observacion_diez_min", "volumen_cinco_min",
            "volumen_diez_min", "suspensibilidad", "dosis_compatibilizador",
        ):
            datos[campo] = _texto(_valor(fila, mapa, campo), 2000)
        for campo in (
            "cantidad_ha", "tasa_aplicacion", "ph_agua", "dureza_agua",
            "ph_agua_corregida", "dureza_agua_corregida", "ph_mezcla",
            "granulometria_300", "granulometria_180", "granulometria_150",
            "granulometria_106", "granulometria_53",
        ):
            datos[campo] = _numero(_valor(fila, mapa, campo))
        datos["agua_corregida"] = _booleano(_valor(fila, mapa, "agua_corregida"))

        for campo, minimo, maximo, etiqueta in [
            ("ph_agua", 0, 14, "pH del agua"),
            ("ph_mezcla", 0, 14, "pH de la mezcla"),
            ("ph_agua_corregida", 0, 14, "pH del agua corregida"),
        ]:
            valor = datos.get(campo)
            if valor is not None and not (minimo <= valor <= maximo):
                lectura.errores.append(
                    ErrorFila(numero_fila, etiqueta, f"Valor fuera de rango ({valor}).")
                )
                datos[campo] = None

        lectura.tratamientos.append(datos)
