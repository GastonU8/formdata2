"""Pruebas de la importación de Excel.

El archivo de prueba se construye acá con las mismas rarezas que trae el archivo
real: encabezados con basura pegada, filas fantasma sin número, decimales con
coma, valores que no son números ("< 0.50"), y una columna `Producto_Final` con
resoluciones equivocadas que el sistema tiene que ignorar.
"""

import pytest
from openpyxl import Workbook

from app import crear_app
from app.extensions import db
from app.models import (
    ArchivoImportado,
    Cliente,
    DetalleMezcla,
    EstadoFlujo,
    EstadoImportacion,
    EstadoNormalizacion,
    EstadoUsuario,
    MapeoProducto,
    OrigenTratamiento,
    Producto,
    ClaveRol,
    Tratamiento,
    Usuario,
)
from app.services import importacion
from app.services import lector_excel as lector
from conftest import aislar_sesion_entre_clientes, rol_de

PASSWORD = "clave-de-prueba-larga"

ENCABEZADOS = [
    "Fecha", "N° de tratamiento", "Tratamiento", "Link", "Año", "Mes", "País", "Lab",
    "Cliente", "NRO CUENTA", "CLIENTE SOFTLAND", "Ubicación", "Clasificación",
    "Cantidad_HA", "Cultivo", "Momento_Aplicacion", "Target", "Tipo_Tratamiento",
    "TasaAplicacion", "PH_Agua", "Dureza_Agua", "AguaCorregida", "PH_AguaCorregida",
    "Dureza_AguaCorregida", "PH_Mezcla", "Observación 5 min", "Volum (ml) 5 MIN",
    "Observación 10 min", "Volum (ml) 10 MIN", "Suspensibilidad (n° inversiones)",
    # Encabezados con basura pegada, tal como vienen en el archivo real:
    "Granulometría (>300 µm)", "Granulometría (>180 µm)22", "Granulometría (>150 µm)2",
    "Granulometría (>106 µm)", "Granulometría (>53 µm)",
    "Dosis de compatibilizador", "Observaciones",
]


def _fila(numero, mezcla, **extra):
    fila = [None] * len(ENCABEZADOS)
    valores = {
        "Fecha": "15/04/2025", "N° de tratamiento": numero, "Tratamiento": mezcla,
        "País": "Argentina", "Lab": "Casa Central", "Cliente": "MSU Agro",
        "NRO CUENTA": "152", "CLIENTE SOFTLAND": "M S U S.A.", "Ubicación": "Villa Cañás",
        "Clasificación": "Grandes Cuentas", "Cultivo": "Soja",
        "Tipo_Tratamiento": "Terrestre", "PH_Agua": 7.96, "Dureza_Agua": 70,
        "AguaCorregida": "SI", "PH_AguaCorregida": 6.77, "PH_Mezcla": "4,04",
        "Observación 5 min": "OK", "Observación 10 min": "PPDO",
        "Volum (ml) 10 MIN": "< 0.50",   # no es un número: es una observación
        "Granulometría (>300 µm)": 0.001, "Granulometría (>180 µm)22": 0.05,
        "Granulometría (>106 µm)": 0.1,
    }
    valores.update(extra)
    for clave, valor in valores.items():
        fila[ENCABEZADOS.index(clave)] = valor
    return fila


def construir_excel(ruta, *, con_hoja_productos=True, numeros_duplicados=False):
    libro = Workbook()

    ws = libro.active
    ws.title = "Tratamientos"
    ws.append(ENCABEZADOS)
    ws.append(_fila(1990, "Sulfato de Amonio + Ligier Grass Bio + Banvel"))
    ws.append(_fila(1991, "Sulfato de Amonio + Atrazina Agroterrum"))
    ws.append(_fila(1991 if numeros_duplicados else 1992, "Voleris + Banvel", PH_Mezcla=99))
    # Filas fantasma: fórmulas arrastradas hasta el fondo de la hoja.
    for _ in range(4):
        fantasma = [None] * len(ENCABEZADOS)
        fantasma[ENCABEZADOS.index("Año")] = 1900
        fantasma[ENCABEZADOS.index("Mes")] = 1
        ws.append(fantasma)

    if con_hoja_productos:
        ws = libro.create_sheet("Productos Normalizados")
        ws.append(["Fecha", "N° de tratamiento", "Producto", "Orden", "Producto_Final", "Id_Producto"])
        # Producto_Final trae resoluciones equivocadas, igual que el archivo real.
        for numero, orden, nombre in [
            (1990, 1, "Sulfato de Amonio"), (1990, 2, "Ligier Grass Bio"), (1990, 3, "Banvel"),
            (1991, 1, "Sulfato de Amonio"), (1991, 2, "Atrazina Agroterrum"),
        ]:
            ws.append(["15/04/2025", numero, nombre, orden, "ADAMA OXALIS 48 SC", 31])

    ws = libro.create_sheet("Mapeo_Productos")
    ws.append(["Producto_Original", "Producto_Final"])
    for original, final in [
        ("Ligier Grass Bio", "LIGIER GRASS"),
        ("Grass", "LIGIER GRASS"),
        ("Banvel", "BANVEL DICAMBA"),
        ("Dédalo", "DEDALO ELITE"),
        ("Dedalo", "DEDALO ELITE"),      # duplicado que solo difiere en el acento
    ]:
        ws.append([original, final])

    ws = libro.create_sheet("Maestro_Productos")
    ws.append(["Id_Producto", "Producto_Final", "N° registro", "Empresa", "Activos",
               "Categoria", "Tipo de Formulación"])
    ws.append([191, "LIGIER GRASS", "39012", "BAHNSA", "Glifosato 60%", "Herbicida", "SL"])
    ws.append([31, "ADAMA OXALIS 48 SC", "38111", "ADAMA", "Flumioxazin 48%", "Herbicida", "SC"])

    libro.save(ruta)
    return str(ruta)


@pytest.fixture
def app():
    aplicacion = aislar_sesion_entre_clientes(crear_app("pruebas"))
    with aplicacion.app_context():
        db.create_all()
        for nombre, email, rol in [
            ("Ana Técnica", "tecnica@bahnsa.com.ar", ClaveRol.TECNICO),
            ("Bruno Comercial", "comercial@bahnsa.com.ar", ClaveRol.COMERCIAL),
        ]:
            usuario = Usuario(nombre=nombre, email=email, rol=rol_de(rol), estado=EstadoUsuario.ACTIVO)
            usuario.definir_password(PASSWORD)
            db.session.add(usuario)
        db.session.commit()
        yield aplicacion
        db.session.remove()
        db.drop_all()


def importar(ruta):
    usuario = db.session.scalar(db.select(Usuario).filter_by(email="tecnica@bahnsa.com.ar"))
    huella, tamano = importacion.calcular_hash(ruta)
    archivo = ArchivoImportado(
        nombre_archivo="prueba.xlsx", hash_sha256=huella, tamano_bytes=tamano,
        usuario_id=usuario.id, estado=EstadoImportacion.PENDIENTE,
    )
    db.session.add(archivo)
    db.session.commit()
    return archivo, importacion.procesar(archivo, ruta)


# ------------------------------------------------------------------- parseo

@pytest.mark.parametrize(
    "texto,esperado",
    [
        ("A + B + C", ["A", "B", "C"]),
        ("  Sulfato de Amonio  +  Banvel  ", ["Sulfato de Amonio", "Banvel"]),
        ("Producto único", ["Producto único"]),
        ("", []),
        (None, []),
    ],
)
def test_la_mezcla_se_parte_por_el_signo_mas(texto, esperado):
    assert lector.partir_mezcla(texto) == esperado


def test_los_encabezados_con_basura_se_reconocen_igual():
    """`Granulometría (>180 µm)22` tiene que seguir siendo la de 180 micras."""
    assert lector._clave_encabezado("N° de tratamiento") == "n de tratamiento"
    assert lector._clave_encabezado("Observación 5 min") == "observacion 5 min"
    mapa = lector._mapa_columnas(tuple(ENCABEZADOS), lector.COLUMNAS_TRATAMIENTO)
    for micras in ("300", "180", "150", "106", "53"):
        assert f"granulometria_{micras}" in mapa
    assert "numero" in mapa and "mezcla_texto" in mapa


def test_los_decimales_con_coma_y_los_no_numeros():
    assert lector._numero("4,04") == 4.04
    assert lector._numero(70) == 70.0
    assert lector._numero("< 0.50") is None      # es una observación, no un valor
    assert lector._numero("") is None
    assert lector._booleano("SI") is True
    assert lector._booleano("no") is False
    assert lector._booleano("tal vez") is None


def test_el_numero_de_tratamiento_tolera_el_formato_de_excel():
    assert lector._numero_tratamiento(1990.0) == "1990"
    assert lector._numero_tratamiento(1990) == "1990"


# ---------------------------------------------------------------- lectura

def test_las_filas_fantasma_se_descartan_sin_error(app, tmp_path):
    lectura = lector.leer(construir_excel(tmp_path / "t.xlsx"))
    assert lectura.filas_leidas == 3
    assert lectura.filas_omitidas == 4
    # Las fantasma no generan ruido: el único error es el pH inválido, que es
    # de una fila real y está puesto a propósito en el archivo de prueba.
    assert [e.columna for e in lectura.errores] == ["pH de la mezcla"]


def test_un_numero_repetido_en_el_mismo_archivo_es_error(app, tmp_path):
    lectura = lector.leer(construir_excel(tmp_path / "t.xlsx", numeros_duplicados=True))
    assert any("ya aparece" in e.mensaje for e in lectura.errores)


def test_sin_hoja_de_productos_se_parte_la_columna_tratamiento(app, tmp_path):
    lectura = lector.leer(construir_excel(tmp_path / "t.xlsx", con_hoja_productos=False))
    primero = lectura.tratamientos[0]
    assert [p["nombre"] for p in primero["productos"]] == [
        "Sulfato de Amonio", "Ligier Grass Bio", "Banvel"
    ]


def test_un_archivo_sin_la_hoja_tratamientos_se_rechaza(app, tmp_path):
    libro = Workbook()
    libro.active.title = "Otra cosa"
    ruta = tmp_path / "malo.xlsx"
    libro.save(ruta)
    with pytest.raises(lector.ArchivoInvalido):
        lector.leer(str(ruta))


def test_el_ph_fuera_de_rango_se_reporta_y_se_descarta(app, tmp_path):
    lectura = lector.leer(construir_excel(tmp_path / "t.xlsx"))
    tercero = next(t for t in lectura.tratamientos if t["numero"] == "1992")
    assert tercero["ph_mezcla"] is None
    assert any("fuera de rango" in e.mensaje for e in lectura.errores)


# ------------------------------------------------------------------- carga

def test_la_importacion_carga_los_tratamientos(app, tmp_path):
    archivo, resumen = importar(construir_excel(tmp_path / "t.xlsx"))
    assert archivo.estado is EstadoImportacion.FINALIZADO_CON_ERRORES  # el pH inválido
    assert resumen.tratamientos_nuevos == 3
    assert resumen.filas_omitidas == 4
    assert db.session.scalar(db.select(db.func.count(Tratamiento.id))) == 3


def test_lo_importado_nace_finalizado_y_no_recorre_el_circuito(app, tmp_path):
    importar(construir_excel(tmp_path / "t.xlsx"))
    for tratamiento in db.session.scalars(db.select(Tratamiento)):
        assert tratamiento.estado_flujo is EstadoFlujo.FINALIZADA
        assert tratamiento.origen is OrigenTratamiento.IMPORTACION
        assert tratamiento.historial == []


def test_los_dos_clientes_se_guardan_por_separado(app, tmp_path):
    importar(construir_excel(tmp_path / "t.xlsx"))
    tratamiento = db.session.scalar(db.select(Tratamiento).filter_by(numero_tratamiento="1990"))
    assert tratamiento.cliente.nombre == "MSU Agro"
    assert tratamiento.cliente_softland.nombre == "M S U S.A."
    assert tratamiento.cliente_softland.numero_cuenta == "152"
    assert db.session.scalar(db.select(db.func.count(Cliente.id))) == 2


def test_las_mediciones_se_cargan_completas(app, tmp_path):
    importar(construir_excel(tmp_path / "t.xlsx"))
    resultado = db.session.scalar(
        db.select(Tratamiento).filter_by(numero_tratamiento="1990")
    ).resultado
    assert float(resultado.ph_mezcla) == 4.04          # decimal con coma
    assert resultado.agua_corregida is True            # "SI"
    assert float(resultado.ph_agua_corregida) == 6.77
    assert float(resultado.granulometria_180) == 0.05  # encabezado con basura
    assert resultado.volumen_diez_min == "< 0.50"      # texto, no número


def test_el_producto_final_equivocado_del_excel_se_ignora(app, tmp_path):
    """La columna del Excel dice ADAMA OXALIS para todo. El sistema no le cree."""
    importar(construir_excel(tmp_path / "t.xlsx"))
    detalle = db.session.scalar(
        db.select(DetalleMezcla).filter_by(nombre_original="Sulfato de Amonio")
    )
    assert detalle.producto is None
    assert detalle.estado_normalizacion is EstadoNormalizacion.PENDIENTE


def test_el_mapeo_del_excel_si_se_respeta(app, tmp_path):
    importar(construir_excel(tmp_path / "t.xlsx"))
    detalle = db.session.scalar(
        db.select(DetalleMezcla).filter_by(nombre_original="Ligier Grass Bio")
    )
    assert detalle.producto.nombre_comercial == "LIGIER GRASS"
    assert detalle.estado_normalizacion is EstadoNormalizacion.CONFIRMADA


def test_los_mapeos_duplicados_por_acento_colapsan(app, tmp_path):
    """«Dédalo» y «Dedalo» son la misma equivalencia, no dos."""
    importar(construir_excel(tmp_path / "t.xlsx"))
    assert db.session.scalar(db.select(db.func.count(MapeoProducto.id))) == 4


def test_el_maestro_del_excel_se_incorpora(app, tmp_path):
    importar(construir_excel(tmp_path / "t.xlsx"))
    producto = db.session.scalar(db.select(Producto).filter_by(nombre_comercial="LIGIER GRASS"))
    assert producto.numero_registro == "39012"
    assert producto.categoria == "Herbicida"
    assert producto.activo == "Glifosato 60%"


def test_reimportar_actualiza_en_lugar_de_duplicar(app, tmp_path):
    ruta = construir_excel(tmp_path / "t.xlsx")
    importar(ruta)
    detalles_primera = db.session.scalar(db.select(db.func.count(DetalleMezcla.id)))

    _, segunda = importar(ruta)
    assert segunda.tratamientos_nuevos == 0
    assert segunda.tratamientos_actualizados == 3
    assert db.session.scalar(db.select(db.func.count(Tratamiento.id))) == 3
    assert db.session.scalar(db.select(db.func.count(DetalleMezcla.id))) == detalles_primera


# -------------------------------------------------------------------- HTTP

def ingresar(app, email):
    cli = app.test_client()
    cli.post("/acceso/", data={"email": email, "password": PASSWORD})
    return cli


def test_solo_el_tecnico_entra_a_importar(app):
    assert ingresar(app, "tecnica@bahnsa.com.ar").get("/importacion/").status_code == 200
    assert ingresar(app, "comercial@bahnsa.com.ar").get("/importacion/").status_code == 403


def test_la_carga_por_formulario_procesa_el_archivo(app, tmp_path):
    ruta = construir_excel(tmp_path / "t.xlsx")
    with open(ruta, "rb") as contenido:
        respuesta = ingresar(app, "tecnica@bahnsa.com.ar").post(
            "/importacion/",
            data={"archivo": (contenido, "historico.xlsx")},
            content_type="multipart/form-data",
        )
    assert respuesta.status_code == 302
    assert db.session.scalar(db.select(db.func.count(Tratamiento.id))) == 3


def test_se_rechaza_un_archivo_que_no_es_excel(app, tmp_path):
    falso = tmp_path / "virus.txt"
    falso.write_text("no soy un excel")
    with open(falso, "rb") as contenido:
        ingresar(app, "tecnica@bahnsa.com.ar").post(
            "/importacion/",
            data={"archivo": (contenido, "virus.txt")},
            content_type="multipart/form-data",
        )
    assert db.session.scalar(db.select(db.func.count(ArchivoImportado.id))) == 0
